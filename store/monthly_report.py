"""
Monthly Attendance Report Generator - eSSL K90 Pro
--------------------------------------------------
Connects to the biometric device and builds a monthly attendance workbook
with two sheets:
  - Summary: one row per staff member, one column per day (P/A), with
    Present/Absent totals - the classic monthly register view.
  - Detail: every individual punch for the month (User Number, Name,
    Date, Time, Punch Type) - the same shape as the daily report.

Reads the same store_config.ini as attendance_report.py, so the device IP
and reports folder only ever need changing in one place.

Usage:
    python monthly_report.py              -> current month, 1st till today
    python monthly_report.py 2026-07      -> a specific month (YYYY-MM),
                                              full month if it's a past month
"""

import calendar
import socket
import sys
from datetime import datetime, date, timedelta
from pathlib import Path

from zk import ZK
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# CONFIGURATION - read from store_config.ini, NOT edited here.
# This file is identical at every store and is updated with `git pull`,
# so each store's own values live in store_config.ini next to it.
# ----------------------------------------------------------------------
from config import settings
settings.require()
DEVICE_IP = settings.get("device", "ip")
DEVICE_PORT = settings.getint("device", "port", 4370)
COMM_PASSWORD = settings.getint("device", "comm_password", 0)
TIMEOUT = settings.getint("device", "timeout", 10)

OUTPUT_FOLDER = settings.getpath("paths", "reports_folder", "attendance_reports")
# ----------------------------------------------------------------------

PUNCH_LABELS = {0: "Check In", 1: "Check Out", 2: "Break Out",
                3: "Break In", 4: "OT In", 5: "OT Out"}

# Excel's standard built-in "Good" (present) / "Bad" (absent) colors
PRESENT_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
PRESENT_FONT = Font(name="Arial", color="006100")
ABSENT_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
ABSENT_FONT = Font(name="Arial", color="9C0006")


def _natural_sort_key(user_id):
    """Numeric user numbers sort in numeric order (2 before 10); anything
    non-numeric falls back to plain string sort, after all numeric ones."""
    try:
        return (0, int(user_id))
    except (TypeError, ValueError):
        return (1, str(user_id))


def probe_tcp_port(ip, port, timeout=5):
    """Raw TCP connection attempt, independent of the ZK protocol - tells
    us whether anything is reachable at ip:port before trying the device
    protocol itself."""
    try:
        with socket.create_connection((ip, port), timeout=timeout):
            return True
    except OSError:
        return False


def fetch_attendance():
    """Connect to the device and return (users, attendance_records).
    Tries TCP first, then UDP, since different devices/firmware prefer
    one or the other."""
    last_error = RuntimeError("no connection attempt was made")
    for use_udp, label in ((False, "TCP"), (True, "UDP")):
        print(f"  attempting connection over {label} ...")
        zk = ZK(DEVICE_IP, port=DEVICE_PORT, timeout=TIMEOUT,
                password=COMM_PASSWORD, force_udp=use_udp, ommit_ping=True)
        conn = None
        try:
            conn = zk.connect()
            conn.disable_device()
            users = conn.get_users()
            attendance = conn.get_attendance()
            print(f"  connected successfully over {label}")
            return users, attendance
        except Exception as e:
            last_error = e
            print(f"  {label} attempt failed: {e}")
        finally:
            if conn:
                try:
                    conn.enable_device()
                except Exception:
                    pass
                try:
                    conn.disconnect()
                except Exception:
                    pass
    raise last_error


def resolve_month_range(month_str=None, today=None):
    """Return (start_date, end_date) for the report.
    No month_str -> current month, capped at today ("till date").
    Explicit month_str ("YYYY-MM") -> that month; if it's the current
    month it's still capped at today, otherwise the full month is used."""
    today = today or date.today()
    if month_str:
        try:
            year, month = (int(p) for p in month_str.split("-"))
        except ValueError:
            print("Month must be in YYYY-MM format, e.g. 2026-07")
            sys.exit(1)
    else:
        year, month = today.year, today.month

    start = date(year, month, 1)
    last_day_num = calendar.monthrange(year, month)[1]
    end = date(year, month, last_day_num)

    if (year, month) == (today.year, today.month):
        end = min(end, today)

    return start, end


def filter_by_date_range(attendance, start_date, end_date):
    """Keep punches between start_date and end_date, inclusive."""
    return [a for a in attendance if start_date <= a.timestamp.date() <= end_date]


def save_monthly_report(users, records, user_lookup, start_date, end_date):
    OUTPUT_FOLDER.mkdir(exist_ok=True)
    filename = OUTPUT_FOLDER / f"attendance_{start_date.isoformat()}_to_{end_date.isoformat()}.xlsx"

    num_days = (end_date - start_date).days + 1
    all_dates = [start_date + timedelta(days=i) for i in range(num_days)]
    users_sorted = sorted(users, key=lambda u: _natural_sort_key(u.user_id))
    punched_days = {(r.user_id, r.timestamp.date()) for r in records}

    wb = Workbook()

    # ---------------- Sheet 1: Summary grid ----------------
    ws1 = wb.active
    ws1.title = "Summary"

    header1 = ["User Number", "Name"] + [str(d.day) for d in all_dates] + ["Present", "Absent"]
    ws1.append(header1)
    for col in range(1, len(header1) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.font = Font(name="Arial", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for u in users_sorted:
        present_count = sum(1 for d in all_dates if (u.user_id, d) in punched_days)
        row_values = [u.user_id, u.name]
        row_values += ["P" if (u.user_id, d) in punched_days else "A" for d in all_dates]
        row_values += [present_count, num_days - present_count]
        ws1.append(row_values)

        r = ws1.max_row
        ws1.cell(row=r, column=1).font = Font(name="Arial")
        ws1.cell(row=r, column=2).font = Font(name="Arial")
        for i, d in enumerate(all_dates):
            cell = ws1.cell(row=r, column=3 + i)
            cell.alignment = Alignment(horizontal="center")
            if (u.user_id, d) in punched_days:
                cell.font = PRESENT_FONT
                cell.fill = PRESENT_FILL
            else:
                cell.font = ABSENT_FONT
                cell.fill = ABSENT_FILL
        ws1.cell(row=r, column=3 + num_days).font = Font(name="Arial", bold=True)
        ws1.cell(row=r, column=4 + num_days).font = Font(name="Arial", bold=True)

    ws1.column_dimensions["A"].width = 12
    ws1.column_dimensions["B"].width = 22
    for i in range(num_days):
        ws1.column_dimensions[get_column_letter(3 + i)].width = 4
    ws1.column_dimensions[get_column_letter(3 + num_days)].width = 9
    ws1.column_dimensions[get_column_letter(4 + num_days)].width = 9
    ws1.freeze_panes = "C2"

    # ---------------- Sheet 2: Detail (flat punch list) ----------------
    ws2 = wb.create_sheet("Detail")
    headers2 = ["User Number", "Name", "Date", "Time", "Punch Type"]
    ws2.append(headers2)
    for col in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.font = Font(name="Arial", bold=True)
        cell.alignment = Alignment(horizontal="center")

    records_sorted = sorted(records, key=lambda a: (_natural_sort_key(a.user_id), a.timestamp))
    for rec in records_sorted:
        punch_value = getattr(rec, "punch", None)
        punch_label = PUNCH_LABELS.get(punch_value,
                                        punch_value if punch_value is not None else "")
        ws2.append([
            rec.user_id,
            user_lookup.get(rec.user_id, ""),
            rec.timestamp.date().isoformat(),
            rec.timestamp.strftime("%H:%M:%S"),
            punch_label,
        ])
        for col in range(1, len(headers2) + 1):
            ws2.cell(row=ws2.max_row, column=col).font = Font(name="Arial")

    for col_idx, header_text in enumerate(headers2, start=1):
        col_letter = get_column_letter(col_idx)
        lengths = [len(str(header_text))]
        for r in range(2, ws2.max_row + 1):
            val = ws2.cell(row=r, column=col_idx).value
            lengths.append(len(str(val)) if val is not None else 0)
        ws2.column_dimensions[col_letter].width = max(lengths) + 4
    ws2.freeze_panes = "A2"

    wb.save(filename)
    return filename


def build_user_lookup(users):
    return {u.user_id: u.name for u in users}


def main():
    month_arg = sys.argv[1] if len(sys.argv) > 1 else None
    start_date, end_date = resolve_month_range(month_arg)

    if not DEVICE_IP:
        print("No device IP set. Fill in [device] ip in store_config.ini.")
        sys.exit(1)

    print(f"Connecting to device at {DEVICE_IP}:{DEVICE_PORT} ...")
    port_open = probe_tcp_port(DEVICE_IP, DEVICE_PORT)
    if port_open:
        print(f"  network check OK - {DEVICE_IP}:{DEVICE_PORT} accepted a TCP connection")
    else:
        print(f"  network check: no plain TCP response from {DEVICE_IP}:{DEVICE_PORT}")

    try:
        users, attendance = fetch_attendance()
    except Exception as e:
        print("\nCould not connect to the device over TCP or UDP.")
        print(f"Last error: {e}\n")
        if not port_open:
            print("The network-level check above also failed - re-check DEVICE_IP,")
            print("that the device is powered on, and that this PC and the device")
            print("are on the same subnet.")
        else:
            print("The device DID respond at the network level, so this is more likely")
            print("a device-side setting (Comm Key, or cloud-only ADMS mode).")
        sys.exit(1)

    print(f"Pulled {len(attendance)} total log entries and "
          f"{len(users)} enrolled users from the device.")

    print(f"Building report for {start_date.isoformat()} to {end_date.isoformat()} "
          f"({(end_date - start_date).days + 1} day(s)) ...")
    ranged_records = filter_by_date_range(attendance, start_date, end_date)
    print(f"{len(ranged_records)} punch(es) found in that range.")

    user_lookup = build_user_lookup(users)
    output_path = save_monthly_report(users, ranged_records, user_lookup, start_date, end_date)
    print(f"Saved report to: {output_path.resolve()}")


if __name__ == "__main__":
    main()
